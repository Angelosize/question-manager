import asyncio
import base64
import threading
import time as _time
from collections import Counter

from fastapi import APIRouter, HTTPException, Depends, Query, Request
from fastapi.responses import HTMLResponse, Response
from typing import List, Dict, Any, Optional
from app.models.paper import PaperCreate, PaperUpdate, Paper
from app.core.data_manager import DataManager, get_data_manager
from app.core.ai_assistant import AIAssistant
from app.models.common import SuccessResponse
import uuid
from datetime import datetime
import re
import io
from urllib.parse import quote

from app.core.config import DATA_DIR

# ---------- 尝试导入 python-docx ----------
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    Document = None
    Inches = Pt = RGBColor = WD_ALIGN_PARAGRAPH = None

router = APIRouter(tags=["papers"])

# ============ 全局常量 ============
SUBJECTS_LIST = ["语文", "数学", "英语", "物理", "化学", "生物", "历史", "地理", "政治"]

# ============ 导入任务状态字典（内存，进程重启即清空） ============
_import_tasks: Dict[str, Dict[str, Any]] = {}
_tasks_lock = threading.Lock()


def _update_task(task_id: str, **kwargs):
    """线程安全地更新任务状态"""
    with _tasks_lock:
        if task_id in _import_tasks:
            _import_tasks[task_id].update(kwargs)


def _get_task(task_id: str) -> Optional[Dict[str, Any]]:
    with _tasks_lock:
        task = _import_tasks.get(task_id)
        return dict(task) if task else None


# ============ 工具函数 ============
def clean_paper_name(name: str) -> str:
    """
    移除卷名开头的年份（如 '2025年'、'2024' 等），只保留纯名称。
    例：'2025年高考化学模拟卷' -> '高考化学模拟卷'
    """
    if not name:
        return name
    cleaned = re.sub(r'^\s*(?:19|20)\d{2}\s*年?\s*', '', name)
    return cleaned.strip()


def _normalize_number(raw: Any) -> Optional[int]:
    """把 '12.' / '12' / 12 / 12.0 统一成 int 12"""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        try:
            return int(raw)
        except (ValueError, OverflowError):
            return None
    if isinstance(raw, str):
        m = re.search(r'(\d+)', raw)
        if m:
            try:
                return int(m.group(1))
            except ValueError:
                return None
    return None


def _extract_page_question_map(valid_pages: List[tuple]) -> Dict[int, int]:
    """
    从多页文本里提取 {题号: 页号(1-based)} 映射。
    valid_pages = [(page_num, text), ...]，page_num 是原始页号（1-based）。
    优先记录题号第一次出现的位置。
    """
    result: Dict[int, int] = {}
    for page_num, text in valid_pages:
        if not text:
            continue
        for m in re.finditer(r'(?:\n|\A)\s*(\d{1,3})\s*[.、．]\s', text):
            try:
                num = int(m.group(1))
            except ValueError:
                continue
            if num not in result:
                result[num] = page_num
    return result


def _fallback_score_by_type(t: str) -> int:
    """按题型给默认分值"""
    if t in ('选择题', '多选题', '填空题'):
        return 3
    if t in ('解答题', '计算题'):
        return 10
    if t in ('实验题', '简答题'):
        return 6
    return 5


def _estimate_duration(question_count: int, subjects: list = None) -> int:
    """
    按题量估测考试时长（分钟）
    """
    if question_count <= 0:
        return 0
    if question_count <= 10:
        per_q = 5.0
    elif question_count <= 20:
        per_q = 4.5
    elif question_count <= 30:
        per_q = 4.0
    else:
        per_q = 3.5
    total = int(question_count * per_q)
    return max(15, round(total / 5) * 5)


def _extract_duration(text: str) -> int:
    """
    从试卷文本中提取考试时长（分钟）
    """
    if not text:
        return 0
    m = re.search(r'\[?\s*考试时长\s*[：:]\s*(\d+)\s*分钟\s*\]?', text)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass
    m = re.search(r'(?:考试时间|时间|时长|限时)\s*[：:]\s*(\d+)\s*分钟', text)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass
    m = re.search(r'(\d+)\s*分钟(?:内|答题|考试|完成)?', text)
    if m:
        try:
            v = int(m.group(1))
            if 15 <= v <= 300:
                return v
        except ValueError:
            pass
    return 0


# ============ ⭐ 学科提取（多级优先） ============
def _extract_subject_from_title(text: str) -> Optional[str]:
    """
    优先级 1：从 [试卷标题：XXX] 里找学科名
    例：[试卷标题：八年级物理晨测(一)] → 物理
    """
    title_match = re.search(r'\[\s*试卷标题\s*[：:]\s*([^\]\n]+?)\s*\]', text)
    if title_match:
        title = title_match.group(1)
        for s in SUBJECTS_LIST:
            if s in title:
                return s
    return None


def _extract_subject_by_frequency(text: str) -> Optional[str]:
    """
    优先级 2：统计全文各学科词出现次数，取最高。
    平票时按 SUBJECTS_LIST 顺序（语文优先）。
    """
    if not text:
        return None
    counts = {s: len(re.findall(re.escape(s), text)) for s in SUBJECTS_LIST}
    candidates = [(s, c) for s, c in counts.items() if c > 0]
    if not candidates:
        return None
    candidates.sort(key=lambda x: (-x[1], SUBJECTS_LIST.index(x[0])))
    return candidates[0][0]


def _extract_subject_by_keyword(text: str) -> Optional[str]:
    """
    优先级 3：关键词顺序硬匹配（旧逻辑，保留兼容）
    """
    if not text:
        return None
    for s in SUBJECTS_LIST:
        if s in text:
            return s
    return None


def _extract_subject_from_split(split_data: List[Dict[str, Any]]) -> Optional[str]:
    """
    优先级 4：AI 单题识别返回的 subject 众数
    """
    if not split_data:
        return None
    counts = Counter()
    for item in split_data:
        s = str(item.get("subject", "")).strip()
        if s and s in SUBJECTS_LIST:
            counts[s] += 1
    if counts:
        return counts.most_common(1)[0][0]
    return None


# ============ 试卷文本处理（支持进度回调 + 页图映射） ============
async def _process_paper_text(
    text: str,
    db: DataManager,
    source_image_bytes: Optional[bytes] = None,
    source_mime: str = "image/jpeg",
    task_id: Optional[str] = None,
    page_images: Optional[List[str]] = None,
    page_question_map: Optional[Dict[int, int]] = None,
) -> Dict[str, Any]:
    """处理试卷文本：提取元信息、AI 分割、创建题目和卷子。"""

    def _progress(stage: str, current: int = 0, total: int = 0):
        if task_id:
            _update_task(task_id, stage=stage, current=current, total=total)

    ai = AIAssistant()
    if not ai.is_available:
        raise HTTPException(status_code=503, detail="AI服务不可用，请检查配置")

    # 提取年份；如果识别不到，使用当前年份
    year_match = re.search(r'(20|19)\d{2}', text)
    year = int(year_match.group()) if year_match else datetime.now().year
    # 提取考试时长
    duration = _extract_duration(text)

    # ⭐ 学科提取（3 级优先，第 4 级在 split 之后）
    subject_found: Optional[str] = None

    # 1. 从标题提取
    subject_found = _extract_subject_from_title(text)
    if subject_found:
        print(f"📌 从标题提取学科: {subject_found}")

    # 2. 从全文词频
    if not subject_found:
        subject_found = _extract_subject_by_frequency(text)
        if subject_found:
            print(f"📌 从全文词频提取学科: {subject_found}")

    # 3. 关键词顺序（旧逻辑兜底）
    if not subject_found:
        subject_found = _extract_subject_by_keyword(text)
        if subject_found:
            print(f"📌 从关键词提取学科: {subject_found}")

    # 提取卷名（优先 AI 标记 → 正则 → 学科兜底）
    paper_name = None
    raw_name = None

    # ⭐ 1. 优先从 AI 标记提取（图片识别会输出 [试卷标题：XXX]）
    title_match = re.search(r'\[\s*试卷标题\s*[：:]\s*([^\]\n]+?)\s*\]', text)
    if title_match:
        candidate = title_match.group(1).strip()
        if candidate and candidate not in ('未标注', '未识别', '无', '未知'):
            raw_name = candidate
            paper_name = clean_paper_name(candidate)
            print(f"📌 从 AI 标记提取卷名: {paper_name}")

    # ⭐ 2. 退一步：从正文里正则找关键词
    if not paper_name:
        name_match = re.search(
            r'([\u4e00-\u9fa5a-zA-Z0-9· ]{2,30}(?:试卷|考试|测试|真题|模拟卷|调研卷|'
            r'基础卷|提高卷|培优卷|单元卷|专题卷|练习卷|作业|练习|练习题|'
            r'月考|期中|期末|联考))',
            text,
        )
        if name_match:
            raw_name = name_match.group(1).strip()
            paper_name = clean_paper_name(raw_name)
            print(f"📌 从正文正则提取卷名: {paper_name}")

    # ⭐ 3. 兜底：学科·章节 / 学科·年份导入
    if not paper_name:
        subject_for_name = subject_found or "未知"
        chapter_match = re.search(r'第[\u4e00-\u9fa5\d]+章\s*[\u4e00-\u9fa5]{2,10}', text)
        if chapter_match:
            chapter = chapter_match.group().strip()
            paper_name = f"{subject_for_name}·{chapter}"
        else:
            paper_name = f"{subject_for_name}·{year}年导入"
        raw_name = paper_name
        print(f"📌 卷名兜底: {paper_name}")

    # ===== 调用 AI（内部已并发）===== → ⭐ 放到线程池
    _progress('splitting', 0, 0)
    split_data = await asyncio.to_thread(
        ai.split_and_recognize, text, progress_callback=_progress
    )
    if not split_data:
        raise HTTPException(status_code=400, detail="未能识别出题目，请检查图片清晰度")

    # ⭐ 4. AI 单题识别众数（兜底）
    if not subject_found:
        subject_found = _extract_subject_from_split(split_data)
        if subject_found:
            print(f"📌 从 AI 识别众数提取学科: {subject_found}")

    # ⭐ 5. 最终兜底
    if not subject_found:
        subject_found = "未知"
        print(f"⚠️ 学科识别失败，使用兜底: 未知")

    print(f"✅ 卷子学科确定为: {subject_found}")

    _progress('saving', 0, len(split_data))

    imported_question_ids = []
    created_questions = []
    for item in split_data:
        try:
            if not item.get("type"):
                item["type"] = "选择题"
            if not item.get("content"):
                item["content"] = "题目内容缺失"
            if not item.get("answer"):
                item["answer"] = ""
            if not item.get("exam_name"):
                item["exam_name"] = paper_name
            if not item.get("year"):
                item["year"] = year
            if not item.get("source"):
                item["source"] = f"{raw_name} {year}年" if year else raw_name
            if not item.get("number"):
                item["number"] = len(created_questions) + 1

            if not item.get("score"):
                item["score"] = _fallback_score_by_type(item.get("type", ""))

            # ⭐ 计算 source_page（题号 → 页号）
            source_page = None
            if page_question_map:
                num_key = _normalize_number(item.get("number"))
                if num_key is not None:
                    source_page = page_question_map.get(num_key)

            question_data = {
                # ⭐ 强制同步卷子科目
                "subject": subject_found,
                "type": item["type"],
                "content": item["content"],
                "answer": item["answer"],
                "difficulty": item.get("difficulty") or "中等",
                "options": item.get("options", {}),
                "explanation": item.get("explanation", ""),
                "knowledge_points": "、".join(item.get("knowledge_points", [])),
                "tags": item.get("knowledge_points", []),
                "source": "AI导入卷子",
            }
            qid = db.add_question(question_data, save=False)

            exam_data = {
                "exam_name": item["exam_name"],
                "year": item["year"],
                "number": item["number"],
                "score": item["score"],
                "source": item["source"],
            }
            if source_page is not None:
                exam_data["source_page"] = source_page
            exam_data = {k: v for k, v in exam_data.items() if v is not None and v != ''}
            if exam_data:
                db.add_exam_mapping(qid, exam_data)

            imported_question_ids.append(qid)
            created_questions.append({
                "question_id": qid,
                "number": item["number"],
                "score": item["score"],
                "source_page": source_page,
            })
        except Exception as e:
            print(f"识别第 {item.get('number', '?')} 题失败: {e}")

    if imported_question_ids:
        db.save_questions()
        print(f"💾 已批量保存 {len(imported_question_ids)} 道题目（学科统一为：{subject_found}）")

    if not imported_question_ids:
        raise HTTPException(status_code=400, detail="所有题目识别失败，请检查图片清晰度")

    # 保存原图（如果有 single source_image_bytes，用于文本导入兼容；多页图片在调用方已保存）
    source_image_name = None
    if source_image_bytes:
        try:
            images_dir = DATA_DIR / "images"
            images_dir.mkdir(parents=True, exist_ok=True)
            ext = "jpg"
            if "png" in source_mime:
                ext = "png"
            elif "webp" in source_mime:
                ext = "webp"
            filename = f"paper_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:6]}.{ext}"
            with open(images_dir / filename, "wb") as f:
                f.write(source_image_bytes)
            source_image_name = filename
            print(f"✅ 试卷原图已保存: {filename}")
        except Exception as e:
            print(f"⚠️ 保存试卷原图失败: {e}")

    # 创建卷子
    total_score = sum(q["score"] for q in created_questions)
    if not duration or duration <= 0:
        duration = _estimate_duration(len(created_questions))
        print(f"⏱️ 按题量估测时长: {duration} 分钟（{len(created_questions)} 题）")
    else:
        print(f"⏱️ AI 识别时长: {duration} 分钟")

    paper_data = {
        "name": paper_name,
        # ⭐ 与题目保持一致的科目
        "subject": subject_found,
        "year": year,
        "duration": duration,
        "total_score": total_score,
        "question_count": len(created_questions),
        "questions": created_questions,
    }
    if page_images:
        clean_page_images = [n for n in page_images if n]
        if clean_page_images:
            paper_data["page_images"] = clean_page_images
            print(f"🖼️ 卷子关联 {len(clean_page_images)} 页原图")
    if source_image_name:
        paper_data["source_image"] = source_image_name
    paper_id = db.add_paper(paper_data)

    return {
        "success": True,
        "paper_id": paper_id,
        "paper_name": paper_name,
        "subject": subject_found,
        "total_questions": len(created_questions),
        "total_score": total_score,
        "duration": duration,
        "source_image": source_image_name,
        "page_images_count": len(page_images) if page_images else 0,
        "message": f"成功导入卷子，共 {len(created_questions)} 道题（学科：{subject_found}）",
    }


# ============ 同步接口：文本导入 ============
@router.post("/import-ai")
async def import_paper_ai(
    request: Dict[str, str],
    db: DataManager = Depends(get_data_manager),
):
    text = request.get("text", "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="文本内容不能为空")
    return await _process_paper_text(text, db)


# ============ 同步接口：图片导入 ============
@router.post("/import-image")
async def import_paper_from_image(
    request: Dict[str, str],
    db: DataManager = Depends(get_data_manager),
):
    image_b64 = request.get("image", "").strip()
    if not image_b64:
        raise HTTPException(status_code=400, detail="图片不能为空")

    mime_type = request.get("mime_type", "image/jpeg")
    try:
        image_bytes = base64.b64decode(image_b64)
    except Exception:
        raise HTTPException(status_code=400, detail="图片 base64 格式错误")
    if len(image_bytes) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="图片过大（>10MB），请压缩后重试")

    ai = AIAssistant()
    if not ai.is_available:
        raise HTTPException(status_code=401, detail="请先配置智谱 API Key")

    # ⭐ 视觉识别放到线程池
    text = await asyncio.to_thread(ai.recognize_paper_image, image_b64, mime_type)
    if not text:
        raise HTTPException(status_code=500, detail="识别失败，请检查图片清晰度或换一张图")

    # 保存单页原图
    images_dir = DATA_DIR / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    ext = "jpg"
    if "png" in mime_type:
        ext = "png"
    elif "webp" in mime_type:
        ext = "webp"
    fname = f"paper_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:6]}_p1.{ext}"
    try:
        with open(images_dir / fname, "wb") as f:
            f.write(image_bytes)
    except Exception as e:
        print(f"⚠️ 保存单页原图失败: {e}")
        fname = None

    return await _process_paper_text(
        text,
        db,
        source_image_bytes=None,
        source_mime=mime_type,
        page_images=[fname] if fname else None,
        page_question_map=_extract_page_question_map([(1, text)]),
    )


# ============ 异步接口：文本导入 ============
@router.post("/import-ai-async")
async def import_paper_ai_async(
    request: Dict[str, str],
    db: DataManager = Depends(get_data_manager),
):
    """异步文本导入，立即返回 task_id"""
    text = request.get("text", "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="文本内容不能为空")

    task_id = uuid.uuid4().hex
    _import_tasks[task_id] = {
        "status": "running",
        "stage": "splitting",
        "current": 0,
        "total": 0,
        "result": None,
        "error": None,
        "started_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    def _run():
        import asyncio
        loop = None
        try:
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            result = loop.run_until_complete(
                _process_paper_text(text, db, task_id=task_id)
            )
            _update_task(task_id, status="done", stage="done", result=result)
        except HTTPException as e:
            _update_task(task_id, status="error", error=str(e.detail))
        except Exception as e:
            import traceback
            traceback.print_exc()
            _update_task(task_id, status="error", error=str(e))
        finally:
            try:
                if loop:
                    loop.close()
            except Exception:
                pass

    threading.Thread(target=_run, daemon=True).start()
    return {"task_id": task_id, "message": "任务已启动"}


# ============ 异步接口：图片导入（支持多页 + 保存所有页） ============
@router.post("/import-image-async")
async def import_paper_from_image_async(
    request: Dict[str, Any],
    db: DataManager = Depends(get_data_manager),
):
    """
    异步图片导入，支持多页。保存所有页原图。
    """
    # 兼容单页 / 多页
    raw_images = request.get("images")
    if not raw_images:
        single = request.get("image", "").strip()
        raw_images = [single] if single else []
    mime_type = request.get("mime_type", "image/jpeg")
    mime_types = request.get("mime_types") or [mime_type] * len(raw_images)

    if not raw_images:
        raise HTTPException(status_code=400, detail="图片不能为空")
    if len(raw_images) > 20:
        raise HTTPException(status_code=400, detail="最多支持 20 页")

    # 解码并校验每页
    image_bytes_list = []
    for idx, b64 in enumerate(raw_images):
        b64 = (b64 or "").strip()
        if not b64:
            raise HTTPException(status_code=400, detail=f"第 {idx+1} 页为空")
        try:
            ib = base64.b64decode(b64)
        except Exception:
            raise HTTPException(status_code=400, detail=f"第 {idx+1} 页 base64 格式错误")
        if len(ib) > 10 * 1024 * 1024:
            raise HTTPException(status_code=413, detail=f"第 {idx+1} 页过大（>10MB）")
        image_bytes_list.append(ib)

    ai = AIAssistant()
    if not ai.is_available:
        raise HTTPException(status_code=401, detail="请先配置智谱 API Key")

    task_id = uuid.uuid4().hex
    _import_tasks[task_id] = {
        "status": "running",
        "stage": "vision",
        "current": 0,
        "total": len(raw_images),
        "result": None,
        "error": None,
        "started_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    def _run():
        import asyncio
        loop = None
        try:
            # ===== 1. 并发识别每一页 =====
            _update_task(task_id, stage="vision", current=0, total=len(raw_images))
            print(f"🖼️ 开始并发识别 {len(raw_images)} 页试卷...")

            page_texts: List[Optional[str]] = [None] * len(raw_images)
            completed = [0]

            def _recognize_page(idx: int):
                try:
                    text = ai.recognize_paper_image(raw_images[idx], mime_types[idx])
                    return text
                except Exception as e:
                    print(f"❌ 第 {idx+1} 页识别失败: {e}")
                    return None

            from concurrent.futures import ThreadPoolExecutor, as_completed
            workers = min(4, len(raw_images))
            with ThreadPoolExecutor(max_workers=workers) as executor:
                future_map = {
                    executor.submit(_recognize_page, i): i
                    for i in range(len(raw_images))
                }
                for future in as_completed(future_map):
                    i = future_map[future]
                    page_texts[i] = future.result()
                    completed[0] += 1
                    _update_task(task_id, stage="vision", current=completed[0], total=len(raw_images))
                    print(f"   ✅ 已识别 {completed[0]}/{len(raw_images)} 页")

            # ===== 2. 保存所有页原图 =====
            images_dir = DATA_DIR / "images"
            images_dir.mkdir(parents=True, exist_ok=True)
            base_ts = datetime.now().strftime('%Y%m%d%H%M%S')
            base_uuid = uuid.uuid4().hex[:6]
            page_image_names: List[Optional[str]] = []
            for idx, ib in enumerate(image_bytes_list, start=1):
                mime = mime_types[idx - 1] if idx - 1 < len(mime_types) else "image/jpeg"
                ext = "jpg"
                if "png" in mime:
                    ext = "png"
                elif "webp" in mime:
                    ext = "webp"
                fname = f"paper_{base_ts}_{base_uuid}_p{idx}.{ext}"
                try:
                    with open(images_dir / fname, "wb") as f:
                        f.write(ib)
                    page_image_names.append(fname)
                except Exception as e:
                    print(f"⚠️ 保存第 {idx} 页失败: {e}")
                    page_image_names.append(None)
            saved_count = len([n for n in page_image_names if n])
            print(f"💾 已保存 {saved_count}/{len(raw_images)} 页原图")

            # ===== 3. 按页合并文本 =====
            valid_pages: List[tuple] = [
                (idx + 1, t) for idx, t in enumerate(page_texts) if t
            ]
            if not valid_pages:
                _update_task(task_id, status="error", error="所有页面均识别失败，请检查图片清晰度")
                return
            merged_text = "\n\n".join(t for _, t in valid_pages)
            print(f"📥 合并后总文本长度: {len(merged_text)} 字符（{len(valid_pages)}/{len(raw_images)} 页成功）")

            # ===== 4. 计算题号 → 页号映射 =====
            page_q_map = _extract_page_question_map(valid_pages)
            if page_q_map:
                preview = dict(sorted(page_q_map.items())[:10])
                suffix = '...' if len(page_q_map) > 10 else ''
                print(f"📌 题号→页号映射: {preview}{suffix}")

            # ===== 5. 走标准文本处理流程 =====
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            result = loop.run_until_complete(
                _process_paper_text(
                    merged_text,
                    db,
                    source_image_bytes=None,
                    source_mime="image/jpeg",
                    task_id=task_id,
                    page_images=page_image_names,
                    page_question_map=page_q_map,
                )
            )

            if isinstance(result, dict):
                result["page_count"] = len(raw_images)
                result["pages_recognized"] = len(valid_pages)

            _update_task(task_id, status="done", stage="done", result=result)
        except HTTPException as e:
            _update_task(task_id, status="error", error=str(e.detail))
        except Exception as e:
            import traceback
            traceback.print_exc()
            _update_task(task_id, status="error", error=str(e))
        finally:
            try:
                if loop:
                    loop.close()
            except Exception:
                pass

    threading.Thread(target=_run, daemon=True).start()
    return {"task_id": task_id, "message": f"任务已启动（{len(raw_images)} 页）"}


# ============ 查询任务状态 ============
@router.get("/import-status/{task_id}")
async def get_import_status(task_id: str):
    task = _get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在或已过期")
    return task


# ============ 卷子页图接口 ============
@router.get("/{paper_id}/page-image/{page_num}")
async def get_paper_page_image(
    paper_id: str,
    page_num: int,
    db: DataManager = Depends(get_data_manager),
):
    """返回卷子第 N 页原图（1-based）"""
    paper = db.get_paper(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="卷子不存在")

    page_images = paper.get("page_images") or []
    if page_num < 1 or page_num > len(page_images):
        raise HTTPException(status_code=404, detail=f"第 {page_num} 页不存在")

    fname = page_images[page_num - 1]
    if not fname:
        raise HTTPException(status_code=404, detail=f"第 {page_num} 页图片缺失")

    image_path = DATA_DIR / "images" / fname
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="图片文件不存在")

    ext = image_path.suffix.lower().lstrip(".")
    media_type = {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "webp": "image/webp",
    }.get(ext, "image/jpeg")

    with open(image_path, "rb") as f:
        content = f.read()
    return Response(content=content, media_type=media_type)


# ============ 已有路由（保持不变） ============
@router.post("/")
async def create_paper(paper: PaperCreate, db: DataManager = Depends(get_data_manager)):
    try:
        paper_data = paper.dict()
        total_score = sum(q.score for q in paper.questions)
        paper_data["total_score"] = total_score
        paper_data["question_count"] = len(paper.questions)
        paper_id = db.add_paper(paper_data)
        return {"id": paper_id, "message": "卷子创建成功"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/")
async def get_all_papers(db: DataManager = Depends(get_data_manager)):
    return db.get_all_papers()


@router.get("/{paper_id}")
async def get_paper(paper_id: str, db: DataManager = Depends(get_data_manager)):
    paper = db.get_paper(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="卷子不存在")
    return paper


@router.put("/{paper_id}")
async def update_paper(paper_id: str, paper: PaperUpdate, db: DataManager = Depends(get_data_manager)):
    update_data = paper.dict(exclude_unset=True)

    if "questions" in update_data and update_data["questions"] is not None:
        questions = update_data["questions"]
        update_data["total_score"] = sum(float(q.get("score", 0) or 0) for q in questions)
        update_data["question_count"] = len(questions)

    success = db.update_paper(paper_id, update_data)
    if not success:
        raise HTTPException(status_code=404, detail="卷子不存在")
    return {"message": "卷子更新成功"}


@router.delete("/{paper_id}")
async def delete_paper(paper_id: str, db: DataManager = Depends(get_data_manager)):
    success = db.delete_paper(paper_id)
    if not success:
        raise HTTPException(status_code=404, detail="卷子不存在")
    return {"message": "卷子删除成功"}


@router.post("/batch-delete", response_model=SuccessResponse)
async def batch_delete_papers(paper_ids: List[str], db: DataManager = Depends(get_data_manager)):
    try:
        if not paper_ids:
            raise HTTPException(status_code=400, detail="请提供要删除的卷子ID列表")
        deleted_count = 0
        failed_ids = []
        for paper_id in paper_ids:
            success = db.delete_paper(paper_id)
            if success:
                deleted_count += 1
            else:
                failed_ids.append(paper_id)
        message = f"成功删除 {deleted_count} 张卷子"
        if failed_ids:
            message += f"，{len(failed_ids)} 张卷子删除失败: {', '.join(failed_ids)}"
        return SuccessResponse(
            message=message,
            data={"deleted_count": deleted_count, "failed_ids": failed_ids},
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"批量删除卷子失败: {str(e)}")


@router.post("/{paper_id}/generate-quiz")
async def generate_quiz_from_paper(paper_id: str, db: DataManager = Depends(get_data_manager)):
    paper = db.get_paper(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="卷子不存在")
    questions = []
    for q in paper["questions"]:
        question = next((q2 for q2 in db.questions if q2["id"] == q["question_id"]), None)
        if question:
            question_copy = question.copy()
            question_copy["number"] = q["number"]
            question_copy["score"] = q["score"]
            question_copy["source_page"] = q.get("source_page")
            questions.append(question_copy)
    return questions


# ============ 导出 HTML ============
@router.get("/{paper_id}/export/html", response_class=HTMLResponse, response_model=None)
async def export_paper_html(
    paper_id: str,
    request: Request,
    db: DataManager = Depends(get_data_manager),
    show_answer: bool = Query(True, description="是否显示答案"),
    show_explanation: bool = Query(True, description="是否显示解析"),
    show_knowledge: bool = Query(True, description="是否显示知识点"),
    show_score: bool = Query(True, description="是否显示分值"),
    answer_position: str = Query("inline", description="答案位置: inline-每题下方, end-末尾集中"),
):
    paper = db.get_paper(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="卷子不存在")

    questions_detail = []
    all_answers = []
    for q in paper.get("questions", []):
        qid = q.get("question_id")
        full_q = next((q2 for q2 in db.questions if q2["id"] == qid), None)
        if full_q:
            exam = db.get_exam_by_question_id(qid)
            full_q_copy = full_q.copy()
            num_raw = q.get("number", exam.get("number") if exam else None)
            if num_raw is not None:
                if isinstance(num_raw, str):
                    match = re.search(r'(\d+)', num_raw)
                    num = int(match.group(1)) if match else None
                else:
                    num = int(num_raw) if isinstance(num_raw, (int, float)) else None
            else:
                num = None
            full_q_copy["number"] = num if num is not None else "?"
            full_q_copy["score"] = q.get("score", exam.get("score") if exam else 0)
            full_q_copy["exam_name"] = exam.get("exam_name") if exam else None
            kp = full_q_copy.get("knowledge_points", "")
            if isinstance(kp, str):
                full_q_copy["knowledge_points"] = [p.strip() for p in kp.split("、") if p.strip()] if kp else []
            elif not isinstance(kp, list):
                full_q_copy["knowledge_points"] = []
            if not isinstance(full_q_copy.get("options"), dict):
                full_q_copy["options"] = {}
            questions_detail.append(full_q_copy)
            all_answers.append({
                "number": full_q_copy["number"],
                "answer": full_q_copy.get("answer", ""),
                "explanation": full_q_copy.get("explanation", "") if show_explanation else None,
            })

    context = {
        "request": request,
        "paper": paper,
        "questions": questions_detail,
        "all_answers": all_answers if answer_position == "end" else [],
        "total_score": paper.get("total_score", 0),
        "question_count": len(questions_detail),
        "generated_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "show_answer": show_answer,
        "show_explanation": show_explanation,
        "show_knowledge": show_knowledge,
        "show_score": show_score,
        "answer_position": answer_position,
    }
    templates = request.app.templates
    template = templates.env.get_template("paper.html")
    html_content = template.render(**context)
    return HTMLResponse(content=html_content)


# ============ 导出 Word ============
@router.get("/{paper_id}/export/docx")
async def export_paper_docx(
    paper_id: str,
    db: DataManager = Depends(get_data_manager),
    show_answer: bool = Query(True, description="是否显示答案"),
    show_explanation: bool = Query(True, description="是否显示解析"),
    show_knowledge: bool = Query(True, description="是否显示知识点"),
    show_score: bool = Query(True, description="是否显示分值"),
):
    if not DOCX_AVAILABLE:
        raise HTTPException(
            status_code=503,
            detail="Word 导出功能未安装，请执行: pip install python-docx",
        )

    paper = db.get_paper(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="卷子不存在")

    questions_detail = []
    for q in paper.get("questions", []):
        qid = q.get("question_id")
        full_q = next((q2 for q2 in db.questions if q2["id"] == qid), None)
        if full_q:
            exam = db.get_exam_by_question_id(qid)
            full_q_copy = full_q.copy()
            num_raw = q.get("number", exam.get("number") if exam else None)
            if num_raw is not None:
                if isinstance(num_raw, str):
                    match = re.search(r'(\d+)', num_raw)
                    num = int(match.group(1)) if match else None
                else:
                    num = int(num_raw) if isinstance(num_raw, (int, float)) else None
            else:
                num = None
            full_q_copy["number"] = num if num is not None else "?"
            full_q_copy["score"] = q.get("score", exam.get("score") if exam else 0)
            full_q_copy["exam_name"] = exam.get("exam_name") if exam else None
            kp = full_q_copy.get("knowledge_points", "")
            if isinstance(kp, str):
                full_q_copy["knowledge_points"] = [p.strip() for p in kp.split("、") if p.strip()] if kp else []
            elif not isinstance(kp, list):
                full_q_copy["knowledge_points"] = []
            if not isinstance(full_q_copy.get("options"), dict):
                full_q_copy["options"] = {}
            questions_detail.append(full_q_copy)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(paper.get("name", "试卷"))
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.name = '宋体'

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_info = f"学科：{paper.get('subject', '未知')}  |  年份：{paper.get('year', '未知')}  |  题量：{len(questions_detail)} 题  |  总分：{paper.get('total_score', 0)} 分"
    run = meta.add_run(meta_info)
    run.font.size = Pt(12)
    run.font.name = '宋体'
    run.font.color.rgb = RGBColor(100, 100, 100)
    doc.add_paragraph()

    for idx, q in enumerate(questions_detail, 1):
        p_question = doc.add_paragraph()
        run = p_question.add_run(f"{q['number']}. ")
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.name = '宋体'
        if show_score:
            run = p_question.add_run(f"（{q['score']} 分）")
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run.font.color.rgb = RGBColor(180, 0, 0)

        p_content = doc.add_paragraph()
        run = p_content.add_run(q.get('content', ''))
        run.font.size = Pt(14)
        run.font.name = '宋体'
        p_content.paragraph_format.left_indent = Inches(0.3)
        p_content.paragraph_format.line_spacing = 1.5

        options = q.get('options', {})
        if options and isinstance(options, dict):
            for letter, text in options.items():
                p_option = doc.add_paragraph()
                p_option.paragraph_format.left_indent = Inches(0.6)
                run = p_option.add_run(f"{letter}. {text}")
                run.font.size = Pt(14)
                run.font.name = '宋体'

        if show_answer:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Inches(0.3)
            p_ans.paragraph_format.space_before = Pt(6)
            run = p_ans.add_run("答案：")
            run.font.bold = True
            run.font.size = Pt(12)
            run = p_ans.add_run(q.get('answer', '未提供'))
            run.font.size = Pt(12)
            run.font.name = '宋体'

            if show_explanation and q.get('explanation'):
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.left_indent = Inches(0.3)
                run = p_exp.add_run("解析：")
                run.font.bold = True
                run.font.size = Pt(12)
                run = p_exp.add_run(q['explanation'])
                run.font.size = Pt(12)
                run.font.name = '宋体'

            if show_knowledge and q.get('knowledge_points'):
                p_kp = doc.add_paragraph()
                p_kp.paragraph_format.left_indent = Inches(0.3)
                run = p_kp.add_run("知识点：")
                run.font.bold = True
                run.font.size = Pt(12)
                run = p_kp.add_run('、'.join(q['knowledge_points']))
                run.font.size = Pt(12)
                run.font.name = '宋体'

        if idx < len(questions_detail):
            doc.add_paragraph('_' * 60).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"由 全能题目管理器 (QTG) 生成  |  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10)
    run.font.name = '宋体'
    run.font.color.rgb = RGBColor(150, 150, 150)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    filename = f"{paper.get('name', '试卷')}.docx"
    encoded_filename = quote(filename)
    return Response(
        content=file_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )